#!/usr/bin/env python3
"""Writes data/KTuner-Maps-Edit-Plan.xlsx: a copy of the owner's digitized map with the
basic-stage edit plan on a first sheet, every cell to change highlighted with a
"before -> after" comment, and each sheet's tab coloured by what to do with it.

The cell list comes from the engine (tools/edit-plan.js), so it matches the app.
Run from the repo root:  python3 tools/annotate-xlsx.py   (needs openpyxl and node)
"""
import json
import os
import subprocess

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'data', 'KTuner-Maps-Digitized.xlsx')
DST = os.path.join(ROOT, 'data', 'KTuner-Maps-Edit-Plan.xlsx')
AUTHOR = 'Civic FE Tune Assist'

plan = json.loads(subprocess.check_output(['node', os.path.join(ROOT, 'tools', 'edit-plan.js')]))
tables = plan['tables']

TAB = {'edit': 'EB6834', 'check': '2A78D6', 'leave': '9AA6B2', 'never': 'C8402F', 'preset': 'C3CCD4', 'info': 'C3CCD4'}
FILL_A = PatternFill('solid', start_color='FDE3CC', end_color='FDE3CC')    # Stage A, required
FILL_B1 = PatternFill('solid', start_color='F9C08F', end_color='F9C08F')   # Lever 1, optional
FILL_B2 = PatternFill('solid', start_color='FFF2D2', end_color='FFF2D2')   # Lever 2, off by default
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical='top')

WHY = {
    'edit': {
        'A': 'Stage A, required after an intake. Paste the corrected 103-value row from the app (Step 3); it is computed from your own log.',
        'B1': 'Lever 1, optional, only after Step 5 passes: full-load columns 7-9 from 3,000 rpm, 11.0 -> 11.5 AFR (11.3 from 5,500 rpm). Same in L and H.',
        'B2': 'Lever 2, optional and OFF by default: +1 psi on the 3,500-5,500 rpm plateau only if a reviewer raises the ceiling above 21 psi. Never below 3,000 rpm.',
    },
    'check': 'Check against logs; do not edit.',
    'leave': 'Leave stock on purpose.',
    'never': 'Never touch: raising knock sensitivity hides knock.',
    'preset': 'Reference curve (KTuner intake preset or factory).',
    'info': 'Partial screenshot, kept to check the digitizing.',
}
REASON = {
    'Final_Boost_Target': 'Axes and exact role not captured in the screenshots.',
    'Boost_By_Gear_Limits': '35.4 psi everywhere: not limiting anything; not a tuning knob in the basic stage.',
    'Cylinder_Fill_Limitation': 'Air-per-cylinder cap: the ECU torque guard that also protects the CVT.',
    'Ignition_Base': 'Knock control finds timing; no dyno means no torque feedback.',
    'Ignition_Max': 'Knock control finds timing; no dyno means no torque feedback.',
    'Ethanol_Ign_Adj': 'Needs an ethanol sensor and KTuner flex-fuel converter. Pump E10 needs no flex setup.',
    'Ethanol_Boost_Target_Adj': 'Needs an ethanol sensor. Pump E10 needs no flex setup.',
    'DI_Fuel_Pressure_Target': 'Values are kPa (18,000 = 180 bar), not psi. Check actual >= 90 % of target in a pull.',
    'Boost_Target_ECO': 'ECO = Stage 1, 18 psi: your hot-day and heavy-traffic mode.',
    'WOT_Exhaust_VTC_Low_Cam': 'Exhaust cam angle at full throttle; shapes spool and scavenging.',
}


def family(name):
    import re
    if re.match(r'^Boost_Target_\d_Normal_', name):
        return 'Boost_Target_Normal'
    if re.match(r'^Boost_Target_\d_ECO_', name):
        return 'Boost_Target_ECO'
    if name.startswith('Knock_Sens_'):
        return 'Knock_Sens'
    if name.startswith('DI_Fuel_Pressure_Target_'):
        return 'DI_Fuel_Pressure_Target'
    return re.sub(r'_(L|H)$', '', name)


def fmt(v):
    return ('%.3f' % v).rstrip('0').rstrip('.') if isinstance(v, float) else str(v)


wb = load_workbook(SRC)
for ws in wb.worksheets:
    t = tables.get(ws.title)
    if not t:
        continue
    role = t['role']
    ws.sheet_properties.tabColor = TAB.get(role, 'C3CCD4')
    note = WHY['edit'][t['stage']] if role == 'edit' else WHY.get(role, '')
    extra = REASON.get(family(ws.title), '')
    if t.get('blindSmoothing') and t['blindSmoothing']['top']:
        b = t['blindSmoothing']
        extra += ' Do not smooth it: a blind 3x3 smooth would add +%s deg at %s rpm (column %s) and raise %d boosted high-load cells by 1 deg or more.' % (
            fmt(b['top']['delta']), b['top']['x'], b['top']['c'], b['raisedHighLoad'])
    ws['A1'].comment = Comment((note + ' ' + extra).strip(), AUTHOR, width=320, height=140)
    curve = ws.max_row <= 2
    fill = {'A': FILL_A, 'B1': FILL_B1, 'B2': FILL_B2}.get(t['stage'])
    for ch in t['changed']:
        # maps: row 1 is the column header and column A the rpm; curves: row 1 is Hz, row 2 the values
        cell = ws.cell(row=2, column=ch['c'] + 1) if curve else ws.cell(row=ch['r'] + 2, column=ch['c'] + 2)
        cell.fill = fill
        cell.font = BOLD
        label = {'A': 'Stage A', 'B1': 'Lever 1 (optional)', 'B2': 'Lever 2 (off by default; only above a 21 psi ceiling)'}[t['stage']]
        cell.comment = Comment('%s: %s -> %s %s' % (label, fmt(ch['from']), fmt(ch['to']), t['unit']), AUTHOR, width=260, height=60)
    if role == 'edit' and t['stage'] == 'A':
        for col in range(1, ws.max_column + 1):
            ws.cell(row=2, column=col).fill = FILL_A
        ws.cell(row=2, column=1).comment = Comment('Stage A: replace this whole row with the corrected row the app computes from your log (Step 3). The app shows it point by point on the Map screen.', AUTHOR, width=300, height=90)

# the plan, first
ps = wb.create_sheet('00_Edit_Plan', 0)
ps.sheet_properties.tabColor = '111B24'
rows = [
    ['Civic FE Tune Assist: the edit plan for your KTuner map (Starter 21 Dual Tune 2 base)'],
    ['One change per flash, one log per change. Cells to change are highlighted on each sheet, with a before -> after comment. Tab colours: orange edit, blue check in logs, grey leave stock, red never touch.'],
    ['Load (column) axes were not captured in the screenshots: check a column\'s real load value in KTuner before you type a change.'],
    [],
    ['Order', 'Sheet (table)', 'Cells', 'Change', 'Only when', 'The log proves it when'],
    ['1 · Stage A (required)', 'MAF_Scaling_Custom (AFM Flow, Custom)', 'All 103 points', 'The corrected row the app computes from your trims and full-throttle error', 'After the intake is fitted; baseline log loaded', 'Trims within +/-5 % at every AFM point driven; full throttle within +/-0.3 AFR of command'],
    ['2 · Lever 1 (optional)', 'WOT_Enrich_L and WOT_Enrich_H', 'Rows 3,000-6,600 rpm x columns 7-9 (27 cells each)', '11.0 -> 11.5 AFR (11.3 from 5,500 rpm)', 'Step 5 passed: a cool-morning and a hot-afternoon log without a Stop', 'Measured within +/-0.3 of command, never leaner than 12.0; no new knock; knock control steady'],
    ['3 · Lever 2 (optional, off by default)', 'Boost_Target_1/2/3_Normal_L and _H', 'Rows 3,500-5,500 rpm x the 21 psi columns (45 cells each)', '21 -> 22 psi', 'A reviewer raises the ceiling above 21 psi, after a clean hot-day log. Never below 3,000 rpm', 'Boost within 1.5 psi of target; knock <= 1 deg; knock control steady; CVT <= 90 C; torque within 3 % of the base map'],
    [],
    ['Leave stock', 'Why'],
]
for name, why in [
    ('Ignition_Base_L/H, Ignition_Max_L/H', REASON['Ignition_Base'] + ' A blind smooth would add timing in the boosted high-load zone.'),
    ('Knock_Sens_* (all four)', 'Never touch: raising sensitivity hides knock, it does not stop it.'),
    ('Boost_Target_*_ECO_*', REASON['Boost_Target_ECO']),
    ('Final_Boost_Target_L/H', REASON['Final_Boost_Target']),
    ('Boost_By_Gear_Limits', REASON['Boost_By_Gear_Limits']),
    ('Cylinder_Fill_Limitation_L/H', REASON['Cylinder_Fill_Limitation']),
    ('Ethanol_Ign_Adj_L/H, Ethanol_Boost_Target_Adj_L/H', REASON['Ethanol_Ign_Adj']),
    ('DI_Fuel_Pressure_Target_0pct/55pct', REASON['DI_Fuel_Pressure_Target']),
    ('WOT_Exhaust_VTC_Low_Cam', REASON['WOT_Exhaust_VTC_Low_Cam']),
    ('MAF_Scaling_Factory, _PRL_Race, _27Won_Race', 'Reference curves. Start from a preset only if you run that exact housing, then correct from your logs.'),
]:
    rows.append([name, why])
rows += [[], ['E10', 'No blanket +4 % fuel. Closed loop trims add it at idle and cruise (2-4 % positive is normal); the Stage A correction carries it; full throttle is proven by the log (trims are off there).']]
for r in rows:
    ps.append(r)
ps['A1'].font = Font(bold=True, size=14)
for c in ps[5]:
    c.font = BOLD
ps['A10'].font = BOLD
ps['B10'].font = BOLD
for col, width in zip('ABCDEF', [34, 40, 34, 36, 40, 46]):
    ps.column_dimensions[col].width = width
for row in ps.iter_rows(min_row=5):
    for c in row:
        c.alignment = WRAP
for r, fill in ((6, FILL_A), (7, FILL_B1), (8, FILL_B2)):
    for c in ps[r]:
        c.fill = fill
wb.active = 0
wb.save(DST)
print('Wrote', os.path.relpath(DST, ROOT))
